# =====================================================================
# Ejabatak Aqrab - AI Inquiry Assistant (Prototype)
#
# Install the required libraries (run once in the terminal):
#   pip install pandas openpyxl scikit-learn pywebview
#
# Run this file to train the model, print the performance report,
# and open the application window.
# =====================================================================
import re
import time
from datetime import datetime
import base64
import mimetypes
import random
import shutil
import os
import webbrowser
from pathlib import Path

import pandas as pd
import webview
from sklearn.model_selection import train_test_split, StratifiedKFold, cross_val_score
from sklearn.pipeline import Pipeline, FeatureUnion
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (accuracy_score, precision_recall_fscore_support,
                             classification_report, confusion_matrix, top_k_accuracy_score)

# ---------- Settings ----------
FILE_PATH = r"C:\Users\rammutairi\Desktop\AI Project.xlsx"
TEXT_COLUMN = "التفاصيل"
TARGET_COLUMN = "معرف الحل"

# Below this confidence, no suggestion is shown ("no match")
NO_MATCH_THRESHOLD = 0.30

# Extra options are shown only if their score is at least this ratio of the top score
CLOSE_RATIO = 0.70
MAX_OPTIONS = 3

# Logo shown in the side panel (optional)
LOGO_PATH = r"C:\Users\rammutairi\Desktop\Logo.png"

# Column in the Solutions sheet holding the service/document link
LINK_COLUMN = "الرابط"
# Folder used for document links written as file names only
DOCUMENTS_FOLDER = Path(FILE_PATH).parent

# Employee ratings log
RATINGS_FILE = Path(FILE_PATH).parent / "ratings_log.xlsx"

# Learning log: every inquiry, prediction, employee action and outcome.
# Rows approved in the review column are added to training on the next run.
LEARNING_FILE = Path(FILE_PATH).parent / "learning_log.xlsx"
LOG_SHEET, SUMMARY_SHEET, RUNS_SHEET = "السجل", "الملخص", "سجل التشغيل"
LEARN_FROM_APPROVED = True

# Cross-validation gives a more reliable accuracy estimate but is slower
RUN_CROSS_VALIDATION = True


# ---------- Text cleaning and model ----------
def clean_arabic_text(text):
    text = str(text)
    text = re.sub(r'[\u0617-\u061A\u064B-\u0652]', '', text)
    text = re.sub(r'[إأآا]', 'ا', text)
    text = text.replace('ى', 'ي').replace('ة', 'ه').replace('ـ', '')
    text = re.sub(r'[^\u0600-\u06FFa-zA-Z0-9\s]', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()


def norm(text):
    """Normalizes Arabic text for matching column names and values."""
    t = str(text)
    t = re.sub(r'[\u200b-\u200f\u202a-\u202e\ufeff\u0640\u064B-\u0652]', '', t)
    t = re.sub(r'[إأآ]', 'ا', t).replace('ة', 'ه').replace('ى', 'ي')
    return re.sub(r'\s+', ' ', t).strip()


# ---------- Learning log ----------
LOG_COLS = ["رقم", "التاريخ", "الاستفسار", "الحل المتوقع", "اسم الحل المتوقع", "نوع الحل", "الثقة",
            "البدائل المعروضة", "تصرف الموظف", "الحل الذي اختاره", "التقييم", "سبب التقييم", "النتيجة",
            "الحل الصحيح (للمراجعة)", "اعتماد للتدريب (نعم/لا)", "حالة التدريب"]
C = {name: i for i, name in enumerate(LOG_COLS)}
REVIEW_COLS = ("الحل الصحيح (للمراجعة)", "اعتماد للتدريب (نعم/لا)")
YES = {"نعم", "yes", "y", "1", "✓", "اي", "ايوه"}
KIND_AR = {"answer": "إجابة", "service": "خدمة", "guide": "مستند"}


def _v(row, col):
    i = C[col]
    val = row[i] if i < len(row) else None
    return "" if val is None else str(val).strip()


def result_of(row):
    """Classifies a logged suggestion as correct, wrong, no match or unknown."""
    top, chosen, behavior = _v(row, "الحل المتوقع"), _v(row, "الحل الذي اختاره"), _v(row, "تصرف الموظف")
    correct, shown = _v(row, "الحل الصحيح (للمراجعة)"), _v(row, "البدائل المعروضة")
    if behavior.startswith("لم نجد"):
        return "لم يقترح"
    if correct:
        if correct == top:
            return "صحيح (مراجَع)"
        return "صحيح ضمن البدائل (مراجَع)" if re.search(rf"\b{re.escape(correct)}\b", shown) else "خطأ (مراجَع)"
    if chosen:
        return "صحيح" if chosen == top else "صحيح ضمن البدائل"
    if "ليس هذا" in behavior:
        return "خطأ"
    return "غير معروف"


def live_stats(rows):
    """Live accuracy based on employee actions and manual review."""
    res = [result_of(r) for r in rows]
    right = sum(r.startswith("صحيح") and "البدائل" not in r for r in res)
    in_alts = sum("ضمن البدائل" in r for r in res)
    wrong = sum(r.startswith("خطأ") for r in res)
    judged = right + in_alts + wrong
    conf = [float(r[C["الثقة"]]) for r in rows if isinstance(r[C["الثقة"]], (int, float))]
    stars = [float(r[C["التقييم"]]) for r in rows if isinstance(r[C["التقييم"]], (int, float))]
    pct = lambda a, b: a / b if b else None
    return {
        "total": len(rows), "no_match": res.count("لم يقترح"), "unknown": res.count("غير معروف"),
        "right": right, "in_alts": in_alts, "wrong": wrong, "judged": judged,
        "accuracy": pct(right, judged), "top3": pct(right + in_alts, judged),
        "avg_conf": pct(sum(conf), len(conf)), "avg_stars": pct(sum(stars), len(stars)), "n_stars": len(stars),
        "reviewed": sum(bool(_v(r, "الحل الصحيح (للمراجعة)")) for r in rows),
        "approved": sum(norm(_v(r, "اعتماد للتدريب (نعم/لا)")).lower() in YES for r in rows),
        "results": res,
    }


def read_log_rows():
    if not LEARNING_FILE.exists():
        return []
    from openpyxl import load_workbook
    try:
        ws = load_workbook(LEARNING_FILE)[LOG_SHEET]
    except KeyError:
        return []
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * (len(LOG_COLS) - len(r))
        if r[C["رقم"]] is not None:
            rows.append(r[:len(LOG_COLS)])
    return rows


def approved_examples(rows, valid_ids):
    """Returns approved log rows as new training examples."""
    texts, labels, status = [], [], {}
    for r in rows:
        rid = r[C["رقم"]]
        approval = norm(_v(r, "اعتماد للتدريب (نعم/لا)")).lower()
        if not approval:
            status[rid] = ""
            continue
        if approval not in YES:
            status[rid] = "غير معتمد"
            continue
        label = _v(r, "الحل الصحيح (للمراجعة)") or _v(r, "الحل الذي اختاره")
        text = clean_arabic_text(_v(r, "الاستفسار"))
        if not label:
            status[rid] = "حدّدي «الحل الصحيح» أولاً"
        elif label not in valid_ids:
            status[rid] = f"المعرّف {label} غير موجود في شيت Solutions"
        elif not text:
            status[rid] = "الاستفسار فاضي"
        else:
            texts.append(text); labels.append(label)
            status[rid] = f"أُضيف للتدريب ({label})"
    return texts, labels, status


class LearningLog:
    """Writes to learning_log.xlsx, reloading the file on each save to keep manual edits."""

    def __init__(self):
        self.pending = []
        self.behavior = {}
        self.top = {}
        rows = read_log_rows()
        self.next_id = max([int(r[C["رقم"]]) for r in rows if str(r[C["رقم"]]).isdigit()] or [0]) + 1

    # Workbook layout
    def _ensure(self, wb):
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        navy = PatternFill("solid", fgColor="003F88")
        amber = PatternFill("solid", fgColor="E8A23A")
        if LOG_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(LOG_SHEET, 0)
            ws.append(LOG_COLS)
            widths = [7, 16, 46, 12, 26, 10, 9, 30, 34, 14, 9, 28, 22, 20, 18, 30]
            for i, (name, w) in enumerate(zip(LOG_COLS, widths), start=1):
                c = ws.cell(row=1, column=i)
                c.fill = amber if name in REVIEW_COLS else navy
                c.font = Font(bold=True, color="1D1D1F" if name in REVIEW_COLS else "FFFFFF")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.column_dimensions[c.column_letter].width = w
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = "D2"
            ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(LOG_COLS)).column_letter}1"
            approve_col = ws.cell(row=1, column=C["اعتماد للتدريب (نعم/لا)"] + 1).column_letter
            dv = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
            dv.add(f"{approve_col}2:{approve_col}20000"); ws.add_data_validation(dv)
            ids = ",".join(solutions["ID"].tolist())
            if len(ids) < 250:
                correct_col = ws.cell(row=1, column=C["الحل الصحيح (للمراجعة)"] + 1).column_letter
                dv2 = DataValidation(type="list", formula1=f'"{ids}"', allow_blank=True)
                dv2.add(f"{correct_col}2:{correct_col}20000"); ws.add_data_validation(dv2)
        for name in (LOG_SHEET, SUMMARY_SHEET, RUNS_SHEET):
            if name not in wb.sheetnames:
                wb.create_sheet(name)
            wb[name].sheet_view.rightToLeft = True
        runs = wb[RUNS_SHEET]
        if runs.cell(1, 1).value is None:
            heads = ["التاريخ", "صفوف التدريب الأصلية", "أمثلة معتمدة أُضيفت", "دقة الاختبار",
                     "التحقق المتقاطع", "الدقة الفعلية من الموظفين حتى الآن", "عدد الاستفسارات المسجلة"]
            for i, (head, w) in enumerate(zip(heads, [16, 18, 18, 14, 26, 26, 22]), start=1):
                c = runs.cell(row=1, column=i, value=head)
                c.fill, c.font = navy, Font(bold=True, color="FFFFFF")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                runs.column_dimensions[c.column_letter].width = w
            runs.row_dimensions[1].height = 32
        if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
            del wb["Sheet"]

    def _row_index(self, ws):
        return {ws.cell(row=i, column=1).value: i for i in range(2, ws.max_row + 1)}

    def _write_summary(self, wb, rows):
        from openpyxl.styles import Font, PatternFill, Alignment
        ws = wb[SUMMARY_SHEET]
        ws.delete_rows(1, ws.max_row)
        s = live_stats(rows)
        f = lambda x: "—" if x is None else f"{x:.2%}"
        ws.append(["ملخص أداء النموذج من استخدام الموظفين", ""])
        ws.append([f"آخر تحديث: {datetime.now():%Y-%m-%d %H:%M}", ""])
        ws.append([])
        lines = [
            ("عدد الاستفسارات المسجلة", s["total"]),
            ("حكمنا عليها (قبول أو رفض أو مراجعة)", s["judged"]),
            ("الدقة الفعلية (الاقتراح الأول صحيح)", f(s["accuracy"])),
            ("الحل الصحيح ضمن الاقتراحات المعروضة", f(s["top3"])),
            ("صحيح", s["right"]), ("صحيح ضمن البدائل", s["in_alts"]), ("خطأ", s["wrong"]),
            ("لم يقترح (ثقة منخفضة)", s["no_match"]),
            ("بدون تفاعل (غير معروف)", s["unknown"]),
            ("متوسط الثقة", f(s["avg_conf"])),
            ("متوسط التقييم", "—" if s["avg_stars"] is None else f"{s['avg_stars']:.2f} / 5  ({s['n_stars']} تقييم)"),
            ("راجعتِ «الحل الصحيح» في", s["reviewed"]),
            ("معتمدة للتدريب", s["approved"]),
        ]
        for line in lines:
            ws.append(list(line))
        ws.append([])
        ws.append(["معرف الحل", "اسم الحل", "مرات الاقتراح", "صحيح", "خطأ", "نسبة الصحة"])
        head = ws.max_row
        per = {}
        for r, res in zip(rows, s["results"]):
            if res == "لم يقترح":
                continue
            d = per.setdefault(_v(r, "الحل المتوقع"), [0, 0, 0])
            d[0] += 1
            d[1] += res.startswith("صحيح") and "البدائل" not in res
            d[2] += res.startswith("خطأ") or "ضمن البدائل" in res
        for sid in sorted(per):
            n, ok, bad = per[sid]
            row = solution_row(sid)
            ws.append([sid, cell(row, "اسم الحل") if row is not None else "", n, ok, bad,
                       f((ok / (ok + bad)) if ok + bad else None)])
        ws.append([])
        ws.append(["كيف تنحسب الدقة الفعلية؟", ""])
        ws.append(["صحيح = الموظف ضغط «نعم» أو فتح الخدمة/المستند للاقتراح الأول", ""])
        ws.append(["صحيح ضمن البدائل = اختار بديلاً من الاقتراحات المعروضة", ""])
        ws.append(["خطأ = ضغط «ليس هذا ما أبحث عنه»", ""])
        ws.append(["إذا عبّيتي «الحل الصحيح» بنفسك، حكمك يلغي تصرف الموظف", ""])
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 44, 30
        for col in "CDEF":
            ws.column_dimensions[col].width = 14
        ws["A1"].font = Font(bold=True, size=14, color="003F88")
        ws["A2"].font = Font(color="6E6E73")
        for c in ws[head]:
            c.fill, c.font = PatternFill("solid", fgColor="003F88"), Font(bold=True, color="FFFFFF")
        for i in range(4, 4 + len(lines)):
            ws.cell(row=i, column=1).font = Font(bold=True)
            ws.cell(row=i, column=2).alignment = Alignment(horizontal="center")
        for i in (6, 7):
            ws.cell(row=i, column=2).font = Font(bold=True, size=12, color="1E7E34")
        ws.cell(row=ws.max_row - 4, column=1).font = Font(bold=True, color="003F88")

    def _refresh(self, ws):
        """Recalculates and colors the result column."""
        from openpyxl.styles import PatternFill
        fills = {"صحيح": "E6F4EA", "صحيح ضمن البدائل": "EAF3FB", "خطأ": "FDE8E8", "لم يقترح": "F1F2F5", "غير معروف": "FFFFFF"}
        rows = []
        for i in range(2, ws.max_row + 1):
            r = [ws.cell(row=i, column=j + 1).value for j in range(len(LOG_COLS))]
            if r[C["رقم"]] is None:
                continue
            res = result_of(r)
            r[C["النتيجة"]] = res
            c = ws.cell(row=i, column=C["النتيجة"] + 1)
            c.value = res
            c.fill = PatternFill("solid", fgColor=fills.get(res.replace(" (مراجَع)", ""), "FFFFFF"))
            rows.append(r)
        return rows

    # Saving
    def flush(self, quiet=False):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment
        try:
            wb = load_workbook(LEARNING_FILE) if LEARNING_FILE.exists() else Workbook()
            self._ensure(wb)
            ws = wb[LOG_SHEET]
            for op in self.pending:
                if op[0] == "new":
                    ws.append(op[1])
                    i = ws.max_row
                    ws.cell(row=i, column=C["الاستفسار"] + 1).alignment = Alignment(wrap_text=True, vertical="top")
                    ws.cell(row=i, column=C["الثقة"] + 1).number_format = "0%"
                elif op[0] == "update":
                    idx = self._row_index(ws)
                    if op[1] in idx:
                        for col, val in op[2].items():
                            ws.cell(row=idx[op[1]], column=C[col] + 1).value = val
                elif op[0] == "status":
                    idx = self._row_index(ws)
                    for rid, text in op[1].items():
                        if rid in idx:
                            ws.cell(row=idx[rid], column=C["حالة التدريب"] + 1).value = text or None
                elif op[0] == "run":
                    runs = wb[RUNS_SHEET]
                    runs.append(op[1])
                    runs.cell(row=runs.max_row, column=4).number_format = "0.00%"
            rows = self._refresh(ws)
            self._write_summary(wb, rows)
            wb.active = 0
            wb.save(LEARNING_FILE)
            self.pending = []
            return rows
        except PermissionError:
            if not quiet:
                print(f"Learning log NOT saved yet: close {LEARNING_FILE.name} in Excel. It will be saved on the next action.")
            return None

    # Events from the interface
    def new_query(self, question, probs, order, shown):
        rid = self.next_id; self.next_id += 1
        top_id = str(model.classes_[order[0]])
        row = solution_row(top_id)
        self.top[rid] = top_id
        if shown:
            alts = "، ".join(f"{o['id']} ({o['conf']:.0%})" for o in shown)
            behavior = "لم يتفاعل"
        else:
            alts = "لم يُعرض شي (ثقة منخفضة)"
            behavior = f"لم نجد إجابة (الثقة أقل من {NO_MATCH_THRESHOLD:.0%})"
        self.behavior[rid] = behavior
        values = [None] * len(LOG_COLS)
        values[C["رقم"]] = rid
        values[C["التاريخ"]] = datetime.now().strftime("%Y-%m-%d %H:%M")
        values[C["الاستفسار"]] = question
        values[C["الحل المتوقع"]] = top_id
        values[C["اسم الحل المتوقع"]] = cell(row, "اسم الحل") if row is not None else ""
        values[C["نوع الحل"]] = KIND_AR.get(kind_of(cell(row, TYPE_COLUMN)), "") if row is not None else ""
        values[C["الثقة"]] = round(float(probs[order[0]]), 4)
        values[C["البدائل المعروضة"]] = alts
        values[C["تصرف الموظف"]] = behavior
        self.pending.append(("new", values))
        self.flush()
        return rid

    def event(self, rid, event, chosen_id="", kind=""):
        if rid is None or rid not in self.behavior:
            return
        top = self.top.get(rid, "")
        alt = " (اختار بديلاً)" if chosen_id and chosen_id != top else ""
        changes = {}
        if event == "yes":
            changes = {"تصرف الموظف": "ضغط «نعم»" + alt, "الحل الذي اختاره": chosen_id}
        elif event == "open":
            what = "فتح الخدمة" if kind == "service" else "فتح المستند"
            changes = {"تصرف الموظف": what + alt, "الحل الذي اختاره": chosen_id}
        elif event == "reject":
            changes = {"تصرف الموظف": "ليس هذا ما أبحث عنه", "الحل الذي اختاره": None}
        elif event == "sent":
            if "أرسل" not in self.behavior[rid]:
                changes = {"تصرف الموظف": self.behavior[rid] + "، ثم أرسل الاستفسار"}
        if "تصرف الموظف" in changes:
            self.behavior[rid] = changes["تصرف الموظف"]
        if changes:
            self.pending.append(("update", rid, changes))
            self.flush()
            print(f"Logged #{rid}: {self.behavior[rid]}")

    def rating(self, rid, stars, reason):
        if rid in self.behavior:
            self.pending.append(("update", rid, {"التقييم": stars, "سبب التقييم": reason or None}))
            self.flush()



def build_model():
    return Pipeline([
        ("features", FeatureUnion([
            ("word_tfidf", TfidfVectorizer(analyzer="word", ngram_range=(1, 2), min_df=2, sublinear_tf=True)),
            ("char_tfidf", TfidfVectorizer(analyzer="char_wb", ngram_range=(2, 5), min_df=2, sublinear_tf=True)),
        ])),
        ("classifier", LogisticRegression(max_iter=3000, class_weight="balanced", random_state=42)),
    ])


# ---------- Load data ----------
training_df = pd.read_excel(FILE_PATH, sheet_name="Training_Data")
solutions = pd.read_excel(FILE_PATH, sheet_name="Solutions")
solutions["ID"] = solutions["ID"].astype(str).str.strip()

raw_rows = len(training_df)
training_df = training_df.dropna(subset=[TEXT_COLUMN, TARGET_COLUMN]).copy()
training_df[TARGET_COLUMN] = training_df[TARGET_COLUMN].astype(str).str.strip()
training_df["clean_text"] = training_df[TEXT_COLUMN].apply(clean_arabic_text)
training_df = training_df[training_df["clean_text"].str.len() > 0]

X, y = training_df["clean_text"], training_df[TARGET_COLUMN]
X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.20, stratify=y, shuffle=True, random_state=42)

# ---------- Add approved examples (training set only, test set unchanged) ----------
try:
    log_rows = read_log_rows()
except PermissionError:
    log_rows = []
    print(f"Could not read {LEARNING_FILE.name}. Close it in Excel to use the approved examples.")
new_texts, new_labels, train_status = approved_examples(log_rows, set(solutions["ID"]))
original_train_rows = len(X_train)
if LEARN_FROM_APPROVED and new_texts:
    X_train = pd.concat([X_train, pd.Series(new_texts)], ignore_index=True)
    y_train = pd.concat([y_train, pd.Series(new_labels)], ignore_index=True)
elif not LEARN_FROM_APPROVED:
    train_status = {k: ("معتمد، لكن LEARN_FROM_APPROVED = False" if v.startswith("أُضيف") else v) for k, v in train_status.items()}
    new_texts = []

# ---------- Train ----------
print("\nTraining model...")
t0 = time.time()
model = build_model()
model.fit(X_train, y_train)
train_seconds = time.time() - t0

# ---------- Evaluate ----------
pred = model.predict(X_test)
proba = model.predict_proba(X_test)
acc = accuracy_score(y_test, pred)
p_m, r_m, f_m, _ = precision_recall_fscore_support(y_test, pred, average="macro", zero_division=0)
p_w, r_w, f_w, _ = precision_recall_fscore_support(y_test, pred, average="weighted", zero_division=0)
n_classes = len(model.classes_)
top3 = top_k_accuracy_score(y_test, proba, k=min(3, n_classes), labels=model.classes_) if n_classes > 3 else 1.0
confidence = proba.max(axis=1)
no_match_rate = float((confidence < NO_MATCH_THRESHOLD).mean())
answered = confidence >= NO_MATCH_THRESHOLD
acc_when_answered = accuracy_score(y_test[answered], pred[answered]) if answered.any() else float("nan")

cv_text = "skipped"
if RUN_CROSS_VALIDATION:
    folds = int(min(5, y.value_counts().min()))
    if folds >= 2:
        print(f"Running {folds}-fold cross-validation...")
        cv = cross_val_score(build_model(), X, y, cv=StratifiedKFold(folds, shuffle=True, random_state=42), scoring="accuracy")
        cv_text = f"{cv.mean():.2%} ± {cv.std():.2%}  ({folds} folds)"
    else:
        cv_text = "skipped (a solution has fewer than 2 examples)"

counts = y.value_counts()
summary = pd.DataFrame([
    ("Rows in Excel", raw_rows),
    ("Rows used after cleaning", len(training_df)),
    ("Training rows (80%)", original_train_rows),
    ("Approved examples added from learning log", len(new_texts)),
    ("Testing rows (20%)", len(X_test)),
    ("Number of solutions (classes)", n_classes),
    ("Examples per solution (min / avg / max)", f"{counts.min()} / {counts.mean():.1f} / {counts.max()}"),
    ("Accuracy", f"{acc:.2%}"),
    ("Top-3 accuracy (right answer in top 3)", f"{top3:.2%}"),
    ("Precision (macro)", f"{p_m:.2%}"),
    ("Recall (macro)", f"{r_m:.2%}"),
    ("F1-score (macro)", f"{f_m:.2%}"),
    ("F1-score (weighted)", f"{f_w:.2%}"),
    ("Cross-validation accuracy", cv_text),
    ("Average confidence on test set", f"{confidence.mean():.2%}"),
    (f"Sent to 'no match' (confidence < {NO_MATCH_THRESHOLD:.0%})", f"{no_match_rate:.2%}"),
    ("Accuracy when a suggestion is shown", f"{acc_when_answered:.2%}"),
    ("Training time", f"{train_seconds:.1f} s"),
], columns=["Metric", "Value"]).set_index("Metric")

pd.set_option("display.width", 200)
pd.set_option("display.max_columns", 50)
line = "=" * 70
print(f"\n{line}\nMODEL PERFORMANCE SUMMARY\n{line}")
print(summary.to_string())

print(f"\n{line}\nPER-SOLUTION REPORT\n{line}")
print(classification_report(y_test, pred, zero_division=0))

print(f"{line}\nCONFUSION MATRIX (rows = actual, columns = predicted)\n{line}")
labels = list(model.classes_)
cm = pd.DataFrame(confusion_matrix(y_test, pred, labels=labels), index=labels, columns=labels)
if n_classes <= 20:
    print(cm.to_string())
else:
    print(f"{n_classes} solutions is too many to print as a grid. Most common mix-ups instead:")
mix = [(a, b, cm.loc[a, b]) for a in labels for b in labels if a != b and cm.loc[a, b] > 0]
mix.sort(key=lambda t: -t[2])
if mix:
    print("\nMost common mix-ups (actual -> predicted : count):")
    for a, b, c in mix[:10]:
        print(f"  {a} -> {b} : {c}")
else:
    print("\nNo mix-ups on the test set.")
print(f"{line}\nOpening the window...\n")


# ---------- Solution types and links ----------
def kind_of(solution_type):
    # Solution type: answer, service or document
    t = norm(solution_type)
    if "خدم" in t:
        return "service"
    if "مستند" in t or "دليل" in t or "سياس" in t or "اجراء" in t:
        return "guide"
    return "answer"


LINK_COLUMN = next((c for c in solutions.columns if norm(c) == norm(LINK_COLUMN)), None) \
    or next((c for c in solutions.columns if "رابط" in norm(c)), LINK_COLUMN)
if LINK_COLUMN not in solutions.columns:
    print(f'WARNING: link column not found. Columns in Solutions: {list(solutions.columns)}')
else:
    n_links = int((solutions[LINK_COLUMN].fillna("").astype(str).str.strip() != "").sum())
    print(f"Link column: '{LINK_COLUMN}' ({n_links} solutions have a link)")

TYPE_COLUMN = next((c for c in solutions.columns if norm(c) == norm("نوع الحل")), "نوع الحل")
if TYPE_COLUMN not in solutions.columns:
    print('WARNING: column "نوع الحل" not found in the Solutions sheet. Every solution will show as an answer.')
else:
    found = solutions[TYPE_COLUMN].fillna("").map(lambda v: (str(v).strip(), kind_of(v))).value_counts()
    print("Solution types found in the Excel file:")
    for (value, kind), n in found.items():
        print(f"  '{value}' -> {kind} ({n})")


# ---------- Documents ----------
def is_url(text):
    return str(text).strip().lower().startswith(("http://", "https://", "www."))


def find_document(text):
    text = str(text).strip().strip('"').strip("'")
    if not text or is_url(text):
        return None
    path = Path(text)
    if not path.is_absolute():
        path = DOCUMENTS_FOLDER / path
    return path if path.is_file() else None


def looks_like_file(text):
    return bool(re.search(r"\.[A-Za-z0-9]{2,5}$", str(text).strip().strip('"'))) and not is_url(text)


def solution_row(solution_id):
    row = solutions[solutions["ID"] == str(solution_id)]
    return None if row.empty else row.iloc[0]


def cell(row, column):
    return "" if column not in row.index or pd.isna(row[column]) else str(row[column]).strip()


# ---------- Update the learning log on each run ----------
LEARN = LearningLog()
live = live_stats(log_rows)
LEARN.pending.append(("status", train_status))
LEARN.pending.append(("run", [datetime.now().strftime("%Y-%m-%d %H:%M"), original_train_rows, len(new_texts),
                              round(acc, 4), cv_text,
                              "—" if live["accuracy"] is None else f"{live['accuracy']:.2%} ({live['judged']} حكم)",
                              live["total"]]))
LEARN.flush()
print(f"\nLearning log: {LEARNING_FILE}")
print(f"  Queries logged so far: {live['total']} | judged: {live['judged']}")
if live["accuracy"] is not None:
    print(f"  Live accuracy from employees: {live['accuracy']:.2%} | right answer among shown options: {live['top3']:.2%}")
print(f"  Approved for training: {live['approved']} | added this run: {len(new_texts)}")
skipped = {k: v for k, v in train_status.items() if v and not v.startswith("أُضيف") and v != "غير معتمد"}
for rid, why in skipped.items():
    print(f"  Row #{rid} not added: {why}")


# ---------- Interface API (called from JavaScript) ----------
WINDOW = None


class Api:
    def ask(self, question):
        probs = model.predict_proba([clean_arabic_text(question)])[0]
        order = probs.argsort()[::-1]
        if probs[order[0]] < NO_MATCH_THRESHOLD:
            print(f"No match: best guess {model.classes_[order[0]]} | confidence {probs[order[0]]:.0%}")
            return {"options": [], "log_id": LEARN.new_query(question, probs, order, [])}
        top = probs[order[0]]
        options = []
        for i in order[:MAX_OPTIONS]:
            if options and probs[i] < top * CLOSE_RATIO:
                break
            sid = str(model.classes_[i])
            row = solution_row(sid)
            if row is None:
                continue
            kind = kind_of(cell(row, TYPE_COLUMN))
            link = cell(row, LINK_COLUMN)
            doc = find_document(link)
            options.append({
                "id": sid,
                "kind": kind,
                "name": cell(row, "اسم الحل"),
                "answer": cell(row, "الحل"),
                "file": doc.name if doc else "",
                "url": link if is_url(link) else "",
                "missing": bool(link and not doc and looks_like_file(link)),
                "has_link": bool(link),
                "conf": float(probs[i]),
            })
        for o in options:
            print(f"Suggested: {o['name']} | confidence {o['conf']:.0%} | type -> {o['kind']} | link: {cell(solution_row(o['id']), LINK_COLUMN) or '(empty)'}")
        return {"options": options, "log_id": LEARN.new_query(question, probs, order, options)}

    def log_event(self, log_id, event, chosen_id="", kind=""):
        LEARN.event(log_id, event, chosen_id, kind)
        return True

    def download(self, solution_id):
        """Saves a copy of the linked document where the employee chooses."""
        row = solution_row(solution_id)
        doc = find_document(cell(row, LINK_COLUMN)) if row is not None else None
        if doc is None:
            return {"ok": False, "msg": "الملف غير موجود في المسار المحدد"}
        dialogs = getattr(webview, "FileDialog", None)
        mode = dialogs.SAVE if dialogs else webview.SAVE_DIALOG
        target = WINDOW.create_file_dialog(mode, save_filename=doc.name)
        if not target:
            return {"ok": False, "cancelled": True}
        target = target if isinstance(target, str) else target[0]
        shutil.copyfile(doc, target)
        return {"ok": True, "saved": str(target)}

    def open_solution_link(self, solution_id):
        """Opens the solution link: URL in the browser, local file with its program."""
        row = solution_row(solution_id)
        link = cell(row, LINK_COLUMN) if row is not None else ""
        if is_url(link):
            self.open_link(link)
            return True
        doc = find_document(link)
        if doc is not None:
            os.startfile(doc)
            return True
        return False

    def open_link(self, url):
        url = (url or "").strip()
        if url.lower().startswith("www."):
            url = "https://" + url
        if url.lower().startswith(("http://", "https://")):
            webbrowser.open(url)
        return True

    def submit_rating(self, stars, reason, solution_name, question="", solution_id="", kind="", log_id=None):
        """Saves the rating to ratings_log.xlsx and prints the average satisfaction."""
        from openpyxl import Workbook, load_workbook
        labels = {"answer": "إجابة", "service": "خدمة", "guide": "مستند"}
        row = [datetime.now().strftime("%Y-%m-%d %H:%M"), question, solution_id, solution_name,
               labels.get(kind, kind), stars, reason]
        try:
            if RATINGS_FILE.exists():
                wb = load_workbook(RATINGS_FILE)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "التقييمات"
                ws.sheet_view.rightToLeft = True
                ws.append(["التاريخ", "الاستفسار", "معرف الحل", "اسم الحل", "نوع الحل", "التقييم", "السبب"])
            ws.append(row)
            wb.save(RATINGS_FILE)
            stars_col = [r[5] for r in ws.iter_rows(min_row=2, values_only=True) if r[5] is not None]
            print(f"Rating saved: {stars}/5 | total ratings: {len(stars_col)} | "
                  f"average satisfaction: {sum(stars_col) / len(stars_col):.2f}/5")
        except PermissionError:
            print(f"Rating NOT saved: close {RATINGS_FILE.name} in Excel and try again.")
        LEARN.rating(log_id, stars, reason)
        return True

    def send_request(self, details, files, log_id=None):
        LEARN.event(log_id, "sent")
        # Demo only: returns a sample request number
        return {"number": f"INQ-{random.randint(100000, 999999)}", "files": len(files)}


# ---------- User interface (HTML / CSS / JavaScript) ----------
HTML = r"""
<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
  :root {
  /* Layout sizes */
    --side-width: clamp(250px, 22vw, 440px);
    --services-width: clamp(150px, 15vw, 290px);
    --item-height: 4.2rem;
    --side-text: #FFFFFF;

    --page: #f7f8fb;
    --card-border: #b9bcc0;
    --field-border: #bdbdbd;
    --grey-field: #f6f6f7;
    --text: #1d1d1f;
    --label: #4a4a4a;
    --hint: #6e6e73;
    --pill: #efefef;
    --btn-bg: #f8f8f8;
    --btn-border: #9a9a9a;
    --red: #e02424;
    --note: #fdf5e8;
    --navy-top: #001f3f;
    --navy-bottom: #003f88;
    --link: #0a4ea3;
    --star: #f5b301;
  }
  html { font-size: clamp(15px, 0.45vw + 11px, 22px); }
  * { box-sizing: border-box; }
  html, body { margin: 0; height: 100%; }
  body {
    font-family: "Segoe UI", Tahoma, sans-serif;
    color: var(--text); background: linear-gradient(180deg, #dcdde2 0%, #eef0f6 55%, #f3f4f9 100%);
    display: grid; grid-template-columns: var(--side-width) 1fr;
  }

  /* Side panel */
  .side { background: linear-gradient(180deg, var(--navy-top) 0%, var(--navy-bottom) 100%); min-height: 100vh; }
  .brand { min-height: 6.5rem; display: flex; align-items: center; }
  .brand .logo { display: block; width: 100%; height: auto; object-fit: contain; }
  .side .tab {
    margin: 1.2rem 1.4rem 0 1.4rem; background: #fff; color: #0d2b57; font-size: 1.02rem; font-weight: 700;
    height: 2.7rem; display: flex; align-items: center; justify-content: flex-start; gap: .5rem;
    padding: 0 .85rem; border-radius: .6rem; white-space: nowrap;
  }

  .side-menu { display: flex; flex-direction: column; padding: .7rem 2.25rem 1rem 1rem; }
  .side-menu span { display: flex; align-items: center; gap: .6rem; color: var(--side-text); font-size: 1rem; padding: .75rem 0; cursor: default; user-select: none; }
  .mi { width: 1rem; height: 1rem; flex: none; fill: none; stroke: currentColor; stroke-width: 1.8; stroke-linecap: round; stroke-linejoin: round; }

  /* Top bar and services column */
  .content { height: 100vh; display: grid; grid-template-rows: auto 1fr; min-width: 0; }
  .topbar {
    background: #fff; border-bottom: 1px solid #c9ccd2; height: 3.4rem;
    display: flex; align-items: center; justify-content: space-between; padding: 0 .9rem 0 4rem;
  }
  .search-deco {
    width: var(--services-width); height: 2.2rem; display: flex; align-items: center; gap: .5rem;
    background: #f1f2f5; border: 1px solid #e1e4ea; border-radius: .6rem; padding: 0 .75rem;
    color: #8a909c; font-size: .8rem; cursor: default; user-select: none;
  }
  .user { display: flex; align-items: center; gap: .55rem; }
  .user-text { display: flex; flex-direction: column; line-height: 1.15; }
  .user-text b { font-size: .85rem; font-weight: 500; color: #222; }
  .user-text span { font-size: .55rem; color: #333; direction: ltr; text-align: left; }
  .user svg { width: 1.65rem; height: 1.65rem; fill: none; stroke: #333; stroke-width: 1.4; }

  .area { display: grid; grid-template-columns: var(--services-width) 1fr; gap: .6rem; padding: .75rem .9rem .75rem 3rem; min-height: 0; }
  .services { background: #fff; border: 1px solid #c9ccd2; border-radius: 1.1rem; overflow: hidden; padding-top: 1.1rem; }
  .services h2 { margin: 0 .75rem 1rem; font-size: 1.15rem; font-weight: 700; color: #0d2b57; display: flex; align-items: center; gap: .5rem; }
  .services a {
    background: #1b9ad6; color: #fff; text-decoration: none; cursor: default; font-size: 1.05rem;
    height: var(--item-height); display: flex; align-items: center; gap: .55rem; padding: 0 .8rem;
  }

  .svc-list span {
    display: flex; align-items: center; gap: .55rem; height: var(--item-height); padding: 0 .8rem;
    background: #fff; color: #111; font-size: 1.05rem; border-top: 1px solid #eef0f3; cursor: default; user-select: none;
  }

  .svc-list .mi { color: #0d2b57; }

  /* Inquiry form */
  .main { overflow-y: auto; min-height: 0; }
  .card {
    min-height: 100%; background: #fff;
    border: 1px solid #c9ccd2; border-radius: 1.1rem; padding: 2.2rem 1.5rem 1.1rem;
  }
  h1 { font-size: 1.25rem; font-weight: 700; margin: 0 .25rem 1.6rem; }

  .field { margin-bottom: 1.6rem; }
  .field > span { display: block; font-size: .75rem; color: var(--label); margin: 0 .25rem .7rem; }
  .star-req { color: var(--red); }
  .box {
    border: 1px solid var(--field-border); border-radius: .9rem; min-height: 3.4rem;
    padding: 0 .75rem; display: flex; align-items: center; font-size: .78rem; color: #444;
  }
  .box.grey { background: var(--grey-field); }
  .field > label > span { display: block; font-size: .75rem; color: var(--label); margin: 0 .25rem .7rem; }
  .select {
    width: 100%; font: inherit; font-size: .78rem; color: #333; background-color: #fff; cursor: pointer; appearance: none; outline: none;
    padding-left: 2.2rem;
    background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' stroke='%23555' stroke-width='2' stroke-linecap='round' stroke-linejoin='round'%3E%3Cpath d='M6 9l6 6 6-6'/%3E%3C/svg%3E");
    background-repeat: no-repeat; background-position: left .8rem center; background-size: 1rem;
  }
  .select:focus { border-color: #6b7f99; box-shadow: 0 0 0 3px #e3e9f2; }
  .select:invalid, .select option[value=""] { color: #777; }
  .select.missing { border-color: var(--red); }
  .box[dir="ltr"] { justify-content: flex-end; }
  textarea {
    width: 100%; min-height: 5.4rem; resize: vertical; border: 1px solid var(--field-border);
    border-radius: .9rem; padding: .65rem .75rem; font: inherit; font-size: .8rem; color: #333; outline: none; line-height: 1.6;
  }
  textarea:focus { border-color: #6b7f99; box-shadow: 0 0 0 3px #e3e9f2; }
  textarea::placeholder { color: #666; }

  /* Suggestion area */
  .suggest-zone { border-top: 1px solid #e3e5e8; border-bottom: 1px solid #e3e5e8; padding: 1.2rem 0 .2rem; margin: .4rem 0 1.2rem; }
  .suggest h2 { font-size: 1.25rem; font-weight: 500; margin: .4rem .25rem .35rem; }
  .suggest .sub { font-size: .875rem; margin: 0 .25rem 1rem; }
  .answer-box {
    border: 1px solid var(--field-border); border-radius: .9rem; min-height: 5.4rem;
    padding: 1.1rem 1.25rem; display: flex; flex-direction: column; justify-content: center;
  }
  .answer-row { display: flex; align-items: center; justify-content: space-between; gap: .9rem; }
  .answer-row .pill { flex: none; }
  .answer-title { font: inherit; font-size: .875rem; text-decoration: underline; text-underline-offset: 4px; color: #111; background: none; border: 0; padding: 0; text-align: right; cursor: default; }
  .answer-title.clickable { cursor: pointer; color: var(--link); }
  .answer-title.clickable:hover { text-decoration-thickness: 2px; }
  .pill { background: var(--pill); color: #6b6b6b; font-size: .65rem; border-radius: 999px; padding: .32rem .9rem; white-space: nowrap; display: inline-flex; align-items: center; gap: .3rem; min-width: 6rem; justify-content: center; }
  .pill svg { width: .8rem; height: .8rem; fill: none; stroke: #6b6b6b; stroke-width: 2; stroke-linecap: round; stroke-linejoin: round; flex: none; }
  .pill .spin { animation: spin 1s linear infinite; }
  @keyframes spin { to { transform: rotate(360deg); } }
  .answer-detail { display: none; margin-top: .9rem; padding-top: .9rem; border-top: 1px dashed #ddd; font-size: .82rem; line-height: 1.9; color: #333; white-space: pre-wrap; word-break: break-word; }
  .answer-detail.open { display: block; }
  .others { display: none; margin-top: .9rem; padding-top: .8rem; border-top: 1px dashed #ddd; }
  .others.show { display: block; }
  .others p { margin: 0 0 .5rem; font-size: .78rem; color: var(--hint); }
  .others ul { list-style: none; margin: 0; padding: 0; display: grid; gap: .45rem; }
  .others li button {
    width: 100%; display: flex; align-items: center; justify-content: space-between; gap: .8rem;
    font: inherit; font-size: .82rem; text-align: right; color: #111; background: #fafafa;
    border: 1px solid #e3e3e3; border-radius: .7rem; padding: .55rem .8rem; cursor: pointer;
  }
  .others li button:hover { background: #f0f3f7; border-color: #c9d2de; }
  .file-link { display: inline-flex; align-items: center; gap: .4rem; font: inherit; font-size: .82rem; color: var(--link); background: none; border: 0; padding: 0; cursor: pointer; text-decoration: underline; }
  .file-link svg { width: 1rem; height: 1rem; fill: none; stroke: currentColor; stroke-width: 2; stroke-linecap: round; stroke-linejoin: round; }
  .status { font-size: .78rem; margin-top: .6rem; color: #1e7e34; display: none; word-break: break-all; }
  .status.show { display: block; }
  .status.bad { color: var(--red); }
  a { color: var(--link); text-decoration: underline; cursor: pointer; }
  a:hover { text-decoration-thickness: 2px; }

  .actions { display: flex; gap: 1rem; margin: 1.6rem 0 1.1rem; }
  .btn {
    font: inherit; font-size: .94rem; height: 2.15rem; padding: 0 1.15rem; border-radius: 1rem;
    border: 1px solid var(--btn-border); background: var(--btn-bg); color: #111; cursor: pointer;
  }
  .btn:hover:not(:disabled) { background: #eef1f5; }
  .btn:disabled { cursor: default; opacity: .6; }
  .btn.a { min-width: 6.25rem; }
  .btn.b { min-width: 9.2rem; }
  .btn:focus-visible, .answer-title:focus-visible, .drop:focus-visible, a:focus-visible, .stars button:focus-visible { outline: 2px solid var(--navy-bottom); outline-offset: 2px; }

  /* Attachments and send */
  .later { display: none; }
  .later.show { display: block; animation: fade .35s ease; }
  @keyframes fade { from { opacity: 0; transform: translateY(-6px); } }
  .later h3 { font-size: .875rem; font-weight: 500; margin: .4rem .25rem .9rem; }
  .drop {
    display: block; border: 1.5px dashed #c4c4c4; margin: 0 .2rem; padding: .9rem .6rem; text-align: center;
    font-size: .82rem; line-height: 1.5; color: #111; cursor: pointer;
  }
  .drop.over { background: #f2f5f9; border-color: #6b7f99; }
  .files { list-style: none; margin: .6rem .2rem 0; padding: 0; display: flex; flex-wrap: wrap; gap: .5rem; }
  .files li { background: var(--pill); border-radius: 999px; font-size: .75rem; padding: .25rem .4rem .25rem .75rem; display: flex; align-items: center; gap: .4rem; }
  .files button { border: 0; background: #d9d9d9; border-radius: 50%; width: 1.15rem; height: 1.15rem; cursor: pointer; font-size: .75rem; line-height: 1; }
  .err { color: var(--red); font-size: .75rem; margin: .5rem .25rem 0; min-height: 1em; }
  .note { display: inline-block; background: var(--note); color: #8a8a8a; font-size: .7rem; padding: .3rem .75rem; margin-top: 1.6rem; }
  .send-row { margin-top: 1.5rem; }
  .send-row .btn { min-width: 7.8rem; }

  .footer { font-size: .7rem; color: #555; margin-top: .6rem; text-align: left; }

  .done { display: none; text-align: center; padding: 1.9rem .6rem; }
  .done.show { display: block; animation: fade .35s ease; }
  .tick { width: 3.5rem; height: 3.5rem; border-radius: 50%; background: #e6f4ea; color: #1e7e34; font-size: 1.9rem; display: grid; place-items: center; margin: 0 auto .75rem; }
  .done p { margin: .25rem 0; }
  .done b { font-size: 1.15rem; }

  /* Rating pop-up */
  .overlay { position: fixed; inset: 0; z-index: 50; display: none; place-items: center; padding: 1.25rem; background: rgba(0, 20, 45, .55); }
  .overlay.show { display: grid; }
  .modal { width: min(100%, 30rem); background: #fff; border-radius: 1rem; padding: 1.6rem 1.6rem 1.3rem; box-shadow: 0 20px 50px rgba(0, 0, 0, .25); animation: pop .3s cubic-bezier(.34, 1.4, .64, 1); }
  @keyframes pop { from { transform: scale(.9); opacity: 0; } }
  .modal h2 { margin: 0 0 .3rem; font-size: 1.2rem; font-weight: 600; }
  .modal .sub { margin: 0 0 1.1rem; font-size: .85rem; color: var(--hint); }
  .stars { display: flex; justify-content: center; gap: .35rem; }
  .stars button { font-size: 2.1rem; line-height: 1; background: none; border: 0; cursor: pointer; color: #d4d4d8; padding: .1rem; transition: transform .12s, color .12s; }
  .stars button.on { color: var(--star); }
  .stars button:hover { transform: scale(1.15); }
  .star-label { text-align: center; font-size: .85rem; color: var(--hint); min-height: 1.3em; margin: .4rem 0 1rem; }
  .modal textarea { min-height: 5rem; }
  .modal .err { text-align: right; }
  .modal .row { display: flex; justify-content: flex-start; margin-top: .9rem; }
  .btn.primary { background: var(--navy-bottom); color: #fff; border-color: var(--navy-bottom); }
  .btn.primary:hover:not(:disabled) { background: #00336f; }


  /* Landing page */
  .landing {
    position: fixed; inset: 0; z-index: 100; overflow: hidden;
    display: flex; align-items: center; justify-content: center; text-align: center; color: #fff;
    background: linear-gradient(160deg, #69BD4B 0%, #006358 48%, #092751 100%);
    transition: opacity .7s ease, visibility .7s ease;
  }
  .landing.gone { opacity: 0; visibility: hidden; }
  .landing .vignette { position: absolute; inset: 0; pointer-events: none;
    background: radial-gradient(ellipse at 50% 45%, rgba(9,39,81,0) 35%, rgba(9,39,81,.38) 100%); }
  .snow { position: absolute; inset: 0; pointer-events: none; }
  .orbit { position: absolute; width: 0; height: 0; animation: orbit linear infinite; }
  .flake { display: block; border-radius: 50%; background: #fff; }
  @keyframes orbit { to { transform: rotate(360deg); } }
  .landing main { position: relative; display: flex; flex-direction: column; align-items: center; padding: 3rem; }
  .landing h1 { margin: 0; font-size: clamp(60px, 9vw, 132px); font-weight: 600; line-height: 1.05; color: #fff; }
  .landing .slogan { margin: 1.4rem 0 0; font-size: clamp(20px, 2.3vw, 32px); font-weight: 300; line-height: 1.5; color: #eaf5ee; }
  .landing .cta {
    margin-top: 4rem; display: inline-flex; align-items: center; gap: .9rem; min-height: 3.75rem; padding: 0 2.5rem;
    border: 1px solid rgba(255,255,255,.45); border-radius: 999px; background: rgba(255,255,255,.16);
    backdrop-filter: blur(16px) saturate(140%); -webkit-backdrop-filter: blur(16px) saturate(140%);
    box-shadow: inset 0 1px 0 rgba(255,255,255,.5), 0 10px 30px rgba(9,39,81,.25);
    color: #fff; font: inherit; font-size: 1.25rem; font-weight: 500; cursor: pointer;
    transition: background .35s ease, border-color .35s ease, box-shadow .35s ease, transform .35s ease;
  }
  .landing .cta:hover { background: rgba(255,255,255,.26); border-color: rgba(255,255,255,.7);
    box-shadow: inset 0 1px 0 rgba(255,255,255,.6), 0 16px 40px rgba(9,39,81,.32); transform: translateY(-2px); }
  .landing .cta svg { width: 1.4rem; height: 1.4rem; fill: none; stroke: #fff; stroke-width: 1.8; stroke-linecap: round; stroke-linejoin: round; transition: transform .35s ease; }
  .landing .cta:hover svg { transform: translateX(-4px); }
  .landing .cta:focus-visible { outline: 2px solid #fff; outline-offset: 6px; }
  .landing .landing-note { margin: 1.1rem 0 0; display: block; background: none; padding: 0; font-size: .95rem; color: #d3e3ef; }

  @media (max-width: 760px) {
    .side .tab { margin: 0; border-radius: 0; }
    .content { height: auto; } .area { grid-template-columns: 1fr; padding: .75rem; }
    body { grid-template-columns: 1fr; }
    .side { min-height: auto; } .side .tab { margin: 0; border-radius: 0; }
    .main { height: auto; padding: 1.2rem; }
  }
  @media (prefers-reduced-motion: reduce) { * { animation: none !important; transition: none !important; } }
</style>
</head>
<body>

<!-- Landing page -->
<section class="landing" id="landing">
  <div class="snow" id="snow" aria-hidden="true"></div>
  <div class="vignette" aria-hidden="true"></div>
  <main>
    <h1>إجابتك أقرب</h1>
    <p class="slogan">ذكاء يختصر الطريق إلى الإجابة</p>
    <button type="button" class="cta" id="enter">
      الانتقال إلى ديوان
      <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M19 12H5"/><path d="M11 6l-6 6 6 6"/></svg>
    </button>
    <p class="landing-note">محاكاة لتجربة الاستفسارات في بيئة ديوان</p>
  </main>
</section>

<aside class="side">
  <div class="brand">__LOGO__</div>
  <div class="tab"><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><circle cx="12" cy="8" r="4"/><path d="M4 20c1.5-3.5 4.5-5 8-5s6.5 1.5 8 5"/></svg><span>الخدمات الذاتية</span></div>
  <div class="side-menu" aria-hidden="true">
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="3" width="7" height="7" rx="1.5"/><rect x="14" y="3" width="7" height="7" rx="1.5"/><rect x="3" y="14" width="7" height="7" rx="1.5"/><rect x="14" y="14" width="7" height="7" rx="1.5"/></svg>التطبيقات</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M3 21h18M5 21V9l7-5 7 5v12M9 21v-6h6v6"/></svg>عن الهيئة</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M3 10v4h4l6 4V6L7 10H3zM17 9a4 4 0 0 1 0 6"/></svg>المركز الإعلامي</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M4 4h5a3 3 0 0 1 3 3v13a2 2 0 0 0-2-2H4zM20 4h-5a3 3 0 0 0-3 3v13a2 2 0 0 1 2-2h6z"/></svg>المكتبة الإلكترونية</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="5" width="18" height="14" rx="2"/><circle cx="9" cy="12" r="2.5"/><path d="M14 10h4M14 14h4"/></svg>هوية الهيئة</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M9 18h6M10 21h4M12 3a6 6 0 0 0-4 10.5c.7.7 1 1.5 1 2.5h6c0-1 .3-1.8 1-2.5A6 6 0 0 0 12 3z"/></svg>مركز المعرفة</span>
    <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M14 3H6a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V9z"/><path d="M14 3v6h6M9 15l2 2 4-4"/></svg>السياسات والاتفاقيات</span>
  </div>
</aside>

<div class="content">
  <header class="topbar">
    <div class="search-deco" aria-hidden="true"><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><circle cx="11" cy="11" r="7"/><path d="M21 21l-4.5-4.5"/></svg><span>عما تبحث؟</span></div>
    <div class="user">
      <div class="user-text"><b>رغد مسير منور المطيري</b><span>Binaa Alkafa'at Trainee</span></div>
      <svg viewBox="0 0 24 24" aria-hidden="true"><circle cx="12" cy="12" r="10.5"/><circle cx="12" cy="9.5" r="3.5"/><path d="M5.5 19c1.4-2.6 3.8-4 6.5-4s5.1 1.4 6.5 4"/></svg>
    </div>
  </header>

  <div class="area">
    <nav class="services" aria-label="الخدمات">
      <h2><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="3" width="18" height="18" rx="3"/><path d="M3 9h18M9 21V9"/></svg>الخدمات</h2>
      <a aria-current="page"><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><circle cx="9" cy="8" r="3.5"/><path d="M2.5 20c1-3.3 3.6-5 6.5-5s5.5 1.7 6.5 5M16 4.5a3.5 3.5 0 0 1 0 7M17.5 15c2.2.5 3.5 2.2 4 5"/></svg>الموارد البشرية</a>
      <div class="svc-list" aria-hidden="true">
        <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="4" width="18" height="12" rx="2"/><path d="M8 20h8M12 16v4"/></svg>تقنية المعلومات</span>
        <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M12 3v18M7 21h10M5 7h14M5 7l-3 6a3 3 0 0 0 6 0zM19 7l-3 6a3 3 0 0 0 6 0z"/></svg>الشؤون القانونية</span>
        <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M3 20h18M6 16l4-5 3 3 5-7"/><path d="M15 7h3v3"/></svg>الاستراتيجية وتيسير التجارة</span>
        <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><path d="M12 3l8 3v6c0 4.5-3.4 8-8 9-4.6-1-8-4.5-8-9V6z"/><path d="M9 12l2 2 4-4"/></svg>أمن المرافق</span>
        <span><svg class="mi" viewBox="0 0 24 24" aria-hidden="true"><ellipse cx="12" cy="5.5" rx="8" ry="2.5"/><path d="M4 5.5v13c0 1.4 3.6 2.5 8 2.5s8-1.1 8-2.5v-13M4 12c0 1.4 3.6 2.5 8 2.5s8-1.1 8-2.5"/></svg>مكتب إدارة البيانات</span>
      </div>
    </nav>

<main class="main">
  <section class="card">
    <h1>الاستفسارات</h1>

    <div class="field">
      <label for="type"><span>نوع الاستفسار<span class="star-req">*</span></span></label>
      <select id="type" class="box select">
        <option value="" disabled selected>اختر نوع الاستفسار</option>
        <option>الرواتب</option>
        <option>خدمات الموظفين</option>
      </select>
    </div>
    <div class="field">
      <label for="topic"><span>الاستفسار<span class="star-req">*</span></span></label>
      <select id="topic" class="box select">
        <option value="" disabled selected>اختر الاستفسار</option>
        <option>المسير الشهري</option>
        <option>طلب شهادة تثبيت راتب</option>
        <option>الحضور والانصراف</option>
      </select>
    </div>

    <div class="field"><span>اسم الموظف</span><div class="box grey">رغد المطيري</div></div>
    <div class="field"><span>الرقم الوظيفي</span><div class="box grey">1211</div></div>
    <div class="field"><span>رقم الهوية / الإقامة</span><div class="box grey">1234567891</div></div>
    <div class="field"><span>الإدارة التابع لها</span><div class="box grey"><bdi>Human Capital Performance</bdi></div></div>

    <div class="field">
      <span>تفاصيل الطلب <span class="star-req">*</span></span>
      <textarea id="details" placeholder="اكتب تفاصيل طلبك هنا"></textarea>
    </div>

    <div id="flow">
      <div class="suggest-zone">
      <div class="suggest">
        <h2>إجابتك أقرب</h2>
        <p class="sub" id="sub">اكتب تفاصيل طلبك وسنقترح عليك إجابة قبل الإرسال</p>
        <div class="answer-box" aria-live="polite">
          <div class="answer-row">
            <button class="answer-title" id="title" tabindex="-1">ستظهر الإجابة المقترحة هنا</button>
            <span class="pill" id="pill"></span>
          </div>
          <div class="answer-detail" id="detail"></div>
          <div class="status" id="status"></div>
          <div class="others" id="others">
            <p>وجدنا إجابة قريبة أخرى، قد يكون قصدك:</p>
            <ul id="others-list"></ul>
          </div>
        </div>
      </div>

      <div class="actions">
        <button class="btn a" id="btnA">ابحث الآن</button>
        <button class="btn b" id="btnB">تخطي الاقتراح</button>
      </div>
      </div>

      <div class="later" id="later">
        <h3>المرفقات</h3>
        <label class="drop" id="drop" tabindex="0">
          <input type="file" id="file" multiple hidden accept=".pdf,.xls,.xlsx,.ppt,.pptx,.doc,.docx,.txt,.rar,.crt,.zip,.p7b,.jpeg,.jpg,.png">
          اسحب ملفك هنا، أو اضغط للاختيار<br>
          يدعم: PDF, EXCEL, POWERPOINT, WORD, TXT, RAR, CRT, ZIP, P7B, JPEG, JPG و PNG<br>
          العدد الأقصى للملفات المسموح برفعها: 3
        </label>
        <ul class="files" id="files"></ul>
        <p class="err" id="err"></p>
        <span class="note">إذا أغلقت هذا الطلب، فستفقد البيانات بشكل دائم</span>
        <div class="send-row"><button class="btn" id="send">إرسال الاستفسار</button></div>
      </div>
    </div>

    <div class="done" id="done">
      <div class="tick">✓</div>
      <p>تم إرسال استفسارك بنجاح</p>
      <b id="reqno"></b>
    </div>

    <div class="footer">كود الخدمة: 1784</div>
  </section>
</main>
  </div>
</div>

<!-- Rating pop-up -->
<div class="overlay" id="rate" role="dialog" aria-modal="true" aria-labelledby="rate-title">
  <div class="modal">
    <h2 id="rate-title">قيّم خدمة إجابتك أقرب</h2>
    <p class="sub">يُسجَّل تقييمك لقياس جودة الإجابات وتحسين النموذج باستمرار</p>
    <div class="stars" id="stars" role="radiogroup" aria-label="التقييم من 1 إلى 5">
      <button role="radio" aria-checked="false" data-v="1" aria-label="1 من 5">★</button>
      <button role="radio" aria-checked="false" data-v="2" aria-label="2 من 5">★</button>
      <button role="radio" aria-checked="false" data-v="3" aria-label="3 من 5">★</button>
      <button role="radio" aria-checked="false" data-v="4" aria-label="4 من 5">★</button>
      <button role="radio" aria-checked="false" data-v="5" aria-label="5 من 5">★</button>
    </div>
    <div class="star-label" id="star-label">اختر عدد النجوم</div>
    <textarea id="reason" placeholder="نأمل توضيح الأسباب في حال كانت الخدمة غير مرضية"></textarea>
    <p class="err" id="rate-err"></p>
    <div class="row"><button class="btn primary" id="rate-send" disabled>إرسال التقييم</button></div>
  </div>
</div>

<script>
  const $ = id => document.getElementById(id);

  /* Icons */
  const I = {
    clock:  '<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/></svg>',
    search: '<svg viewBox="0 0 24 24" class="spin"><path d="M21 12a9 9 0 1 1-6.2-8.6"/></svg>',
    doc:    '<svg viewBox="0 0 24 24"><path d="M14 3H7a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h10a2 2 0 0 0 2-2V8z"/><path d="M14 3v5h5M9 13h6M9 17h6"/></svg>',
    book:   '<svg viewBox="0 0 24 24"><path d="M4 5a2 2 0 0 1 2-2h13v16H6a2 2 0 0 0-2 2z"/><path d="M4 19V5M8 7h7"/></svg>',
    brief:  '<svg viewBox="0 0 24 24"><rect x="3" y="7" width="18" height="13" rx="2"/><path d="M9 7V5a2 2 0 0 1 2-2h2a2 2 0 0 1 2 2v2M3 13h18"/></svg>',
    none:   '<svg viewBox="0 0 24 24"><circle cx="11" cy="11" r="7"/><path d="M21 21l-4.3-4.3M8.5 8.5l5 5M13.5 8.5l-5 5"/></svg>',
    down:   '<svg viewBox="0 0 24 24"><path d="M12 4v11m0 0l-4-4m4 4l4-4M5 19h14"/></svg>',
    check:  '<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="9"/><path d="M8 12.5l2.5 2.5L16 9.5"/></svg>',
  };

  /* Links inside answers */
  function escapeHtml(s) { return String(s).replace(/[&<>"']/g, c => ({ "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;" }[c])); }
  function linkify(text) {
    return escapeHtml(text).replace(/(https?:\/\/[^\s<]+|www\.[^\s<]+)/gi, url => {
      const clean = url.replace(/[.,،؛)]+$/, "");
      const tail = url.slice(clean.length);
      const href = clean.startsWith("http") ? clean : "https://" + clean;
      return `<a href="${href}" dir="ltr">${clean}</a>${tail}`;
    });
  }
  const isLink = s => /^(https?:\/\/|www\.)/i.test((s || "").trim());
  document.addEventListener("click", e => {
    const a = e.target.closest("a[href]");
    if (!a) return;
    e.preventDefault();
    window.pywebview.api.open_link(a.getAttribute("href"));
  });

  /* Suggestion states */
  let current = null, last = null, files = [], logId = null;
  const logEvent = ev => { if (logId != null) window.pywebview.api.log_event(logId, ev, last ? last.id : "", last ? last.kind : ""); };

  function setState(state, d = {}) {
    const title = $("title"), detail = $("detail");
    detail.classList.remove("open"); detail.innerHTML = "";
    $("status").className = "status"; $("status").textContent = "";
    if (!["answer", "guide", "service"].includes(state)) { $("others").classList.remove("show"); }
    title.classList.remove("clickable"); title.tabIndex = -1;
    $("btnA").disabled = false; $("btnB").disabled = false;

    const S = {
      idle:      { sub: "اكتب تفاصيل طلبك وسنقترح عليك إجابة قبل الإرسال", title: "ستظهر الإجابة المقترحة هنا", pill: I.clock + "بانتظار الطلب", a: "ابحث الآن", b: "تخطي الاقتراح" },
      searching: { sub: "نبحث عن إجابة مناسبة لطلبك...", title: "جارٍ البحث...", pill: I.search + "جارٍ البحث", a: "انتظر...", b: "تخطي الاقتراح" },
      answer:    { sub: "هل تجيب هذه الإجابة على سؤالك؟", pill: I.doc + "إجابة فورية", a: "نعم", b: "ليس هذا ما أبحث عنه" },
      guide:     { sub: `هل تقصد الاستفسار عن ${d.name || ""}؟`, pill: I.book + "سياسة / إجراء", a: "الانتقال إلى المستند", b: "ليس هذا ما أبحث عنه" },
      service:   { sub: `هل تقصد ${d.name || ""}؟`, pill: I.brief + "خدمة / طلب", a: "الانتقال إلى الخدمة", b: "ليس هذا ما أبحث عنه" },
      none:      { sub: "لم نجد إجابة مناسبة لطلبك", title: "أكمل استفسارك وأرفق ما يلزم، ثم أرسله", pill: I.none + "لا توجد نتائج", a: "إعادة البحث", b: "متابعة الاستفسار" },
      helped:    { sub: "شكراً لتقييمك! يسعدنا أن الإجابة أفادتك", pill: I.check + "تم الحل", a: "استفسار جديد", b: "" },
    }[state];

    current = state;
    $("sub").textContent = S.sub;
    title.textContent = S.title || d.name || "";
    $("pill").innerHTML = S.pill;
    $("btnA").textContent = S.a;
    $("btnB").textContent = S.b;
    $("btnB").style.display = S.b ? "" : "none";
    if (state === "searching") $("btnA").disabled = true;

    if (["answer", "guide", "service", "helped"].includes(state)) {
      if (d.answer) detail.innerHTML = linkify(d.answer);
      if (d.file && state !== "helped") {
        const wrap = document.createElement("div");
        wrap.style.marginTop = d.answer ? ".6rem" : "0";
        wrap.innerHTML = `<button class="file-link">${I.down}<span dir="auto"></span></button>`;
        wrap.querySelector("span").textContent = d.file;
        wrap.querySelector("button").onclick = downloadDoc;
        detail.appendChild(wrap);
      } else if (d.missing && state !== "helped") {
        const miss = document.createElement("div");
        miss.style.cssText = "margin-top:.6rem;color:#8a8a8a;font-size:.78rem";
        miss.textContent = "المستند غير متوفر حالياً";
        detail.appendChild(miss);
      }
      if (detail.childNodes.length) detail.classList.add("open");
      if (d.url) { title.classList.add("clickable"); title.tabIndex = 0; }
    }
  }

  async function search() {
    const text = $("details").value.trim();
    if (text.length < 3) { $("details").focus(); $("sub").textContent = "اكتب تفاصيل طلبك أولاً، ثم اضغط «ابحث الآن»"; return; }
    hideLater();
    setState("searching");
    const [d] = await Promise.all([window.pywebview.api.ask(text), new Promise(r => setTimeout(r, 900))]);
    options = d.options || [];
    logId = d.log_id ?? null;
    if (!options.length) return setState("none");
    choose(0);
  }

  let options = [];
  function choose(i) {
    last = options[i];
    setState(last.kind, last);
    const rest = options.map((o, j) => ({ o, j })).filter(x => x.j !== i);
    const list = $("others-list");
    list.innerHTML = "";
    rest.forEach(({ o, j }) => {
      const li = document.createElement("li");
      const b = document.createElement("button");
      const name = document.createElement("span"); name.textContent = o.name;
      const pill = document.createElement("span"); pill.className = "pill";
      pill.innerHTML = { answer: I.doc + "إجابة فورية", guide: I.book + "سياسة / إجراء", service: I.brief + "خدمة / طلب" }[o.kind];
      b.append(name, pill);
      b.onclick = () => choose(j);
      li.appendChild(b); list.appendChild(li);
    });
    $("others").classList.toggle("show", rest.length > 0);
  }

  async function downloadDoc() {
    if (!last || !last.file) return;
    const r = await window.pywebview.api.download(last.id);
    const s = $("status");
    if (r.cancelled) return;
    s.className = "status show" + (r.ok ? "" : " bad");
    s.textContent = r.ok ? "تم حفظ المستند في: " : r.msg;
    if (r.ok) { const path = document.createElement("bdi"); path.dir = "ltr"; path.textContent = r.saved; s.appendChild(path); }
  }

  async function openSolutionLink() {
    const ok = last && last.has_link ? await window.pywebview.api.open_solution_link(last.id) : false;
    if (ok) { logEvent("open"); openRating(); return; }
    if (!ok) { const s = $("status"); s.className = "status show bad"; s.textContent = last && last.has_link ? "الملف أو الرابط غير موجود في المسار المحدد" : "لم يُضف رابط لهذا الحل بعد في عمود «الرابط»"; }
  }

  function showLater() { $("later").classList.add("show"); $("later").scrollIntoView({ behavior: "smooth", block: "start" }); }
  function hideLater() { $("later").classList.remove("show"); }
  function openAnswerLink() {
    const u = last && (last.url || (isLink(last.answer) ? last.answer.trim() : ""));
    if (u) window.pywebview.api.open_link(u);
  }

  $("btnA").onclick = () => {
    if (current === "idle" || current === "none") return search();
    if (current === "answer") { logEvent("yes"); return openRating(); }
    if (current === "guide" || current === "service") return openSolutionLink();
    if (current === "helped") { $("details").value = ""; hideLater(); setState("idle"); $("details").focus(); }
  };
  $("btnB").onclick = () => {
    if (current === "searching") return;
    if (["answer", "guide", "service"].includes(current)) logEvent("reject");
    showLater();
  };
  $("title").onclick = () => { if ($("title").classList.contains("clickable")) openAnswerLink(); };

  $("details").addEventListener("input", () => {
    if (current !== "idle" && current !== "searching") { setState("idle"); hideLater(); }
  });
  $("details").addEventListener("keydown", e => { if (e.key === "Enter" && e.ctrlKey) search(); });

  /* Rating */
  const LABELS = ["", "غير مرضية إطلاقاً", "غير مرضية", "مقبولة", "جيدة", "ممتازة"];
  let stars = 0;
  function paintStars(n) { [...$("stars").children].forEach(b => b.classList.toggle("on", +b.dataset.v <= n)); }
  function openRating() {
    stars = 0; paintStars(0); $("reason").value = ""; $("rate-err").textContent = "";
    $("star-label").textContent = "اختر عدد النجوم"; $("rate-send").disabled = true;
    $("rate").classList.add("show");
    $("stars").children[0].focus();
  }
  [...$("stars").children].forEach(b => {
    b.addEventListener("mouseenter", () => paintStars(+b.dataset.v));
    b.addEventListener("click", () => {
      stars = +b.dataset.v;
      [...$("stars").children].forEach(x => x.setAttribute("aria-checked", +x.dataset.v === stars));
      $("star-label").textContent = LABELS[stars];
      $("rate-send").disabled = false; $("rate-err").textContent = "";
      paintStars(stars);
    });
  });
  $("stars").addEventListener("mouseleave", () => paintStars(stars));
  $("stars").addEventListener("keydown", e => {
    const k = { ArrowLeft: 1, ArrowUp: 1, ArrowRight: -1, ArrowDown: -1 }[e.key];
    if (!k) return;
    e.preventDefault();
    const v = Math.min(5, Math.max(1, (stars || 0) + k));
    $("stars").children[v - 1].click(); $("stars").children[v - 1].focus();
  });
  document.addEventListener("keydown", e => { if (e.key === "Escape" && $("rate").classList.contains("show")) e.preventDefault(); });
  $("rate-send").onclick = async () => {
    const reason = $("reason").value.trim();
    if (!stars) { $("rate-err").textContent = "اختر عدد النجوم أولاً"; return; }
    if (stars <= 2 && !reason) { $("rate-err").textContent = "نأمل توضيح الأسباب لأن التقييم غير مرضٍ"; $("reason").focus(); return; }
    $("rate-send").disabled = true;
    await window.pywebview.api.submit_rating(stars, reason, last ? last.name : "", $("details").value.trim(), last ? last.id : "", last ? last.kind : "", logId);
    $("rate").classList.remove("show");
    setState("helped", last);
  };

  /* Attachments */
  const OK = /\.(pdf|xlsx?|pptx?|docx?|txt|rar|crt|zip|p7b|jpe?g|png)$/i;
  function addFiles(list) {
    $("err").textContent = "";
    for (const f of list) {
      if (!OK.test(f.name)) { $("err").textContent = `نوع الملف غير مدعوم: ${f.name}`; continue; }
      if (files.length >= 3) { $("err").textContent = "العدد الأقصى للملفات هو 3"; break; }
      if (!files.includes(f.name)) files.push(f.name);
    }
    renderFiles();
  }
  function renderFiles() {
    $("files").innerHTML = "";
    files.forEach((n, i) => {
      const li = document.createElement("li");
      li.textContent = "📎 " + n;
      const x = document.createElement("button");
      x.textContent = "×"; x.setAttribute("aria-label", "حذف " + n);
      x.onclick = () => { files.splice(i, 1); renderFiles(); };
      li.appendChild(x); $("files").appendChild(li);
    });
  }
  $("file").onchange = () => { addFiles($("file").files); $("file").value = ""; };
  $("drop").addEventListener("keydown", e => { if (e.key === "Enter" || e.key === " ") { e.preventDefault(); $("file").click(); } });
  ["dragenter", "dragover"].forEach(ev => $("drop").addEventListener(ev, e => { e.preventDefault(); $("drop").classList.add("over"); }));
  ["dragleave", "drop"].forEach(ev => $("drop").addEventListener(ev, e => { e.preventDefault(); $("drop").classList.remove("over"); }));
  $("drop").addEventListener("drop", e => addFiles(e.dataTransfer.files));

  /* Send */
  $("send").onclick = async () => {
    const text = $("details").value.trim();
    for (const id of ["type", "topic"]) {
      if (!$(id).value) { $(id).classList.add("missing"); $("err").textContent = "يرجى اختيار نوع الاستفسار والاستفسار"; $(id).focus(); return; }
    }
    if (!text) { $("err").textContent = "تفاصيل الطلب مطلوبة"; $("details").focus(); return; }
    $("send").disabled = true;
    const r = await window.pywebview.api.send_request(text, files, logId);
    $("flow").style.display = "none";
    $("details").readOnly = true;
    $("reqno").textContent = "رقم الطلب: " + r.number;
    $("done").classList.add("show");
  };

  ["type", "topic"].forEach(id => $(id).addEventListener("change", () => { $(id).classList.remove("missing"); $("err").textContent = ""; }));

  /* Landing page */
  (function () {
    const snow = $("snow");
    let seed = 7;
    const rnd = () => { seed = (seed * 9301 + 49297) % 233280; return seed / 233280; };
    for (let i = 0; i < 70; i++) {
      const size = 2 + rnd() * 4, radius = 18 + rnd() * 70, duration = 28 + rnd() * 44;
      const o = document.createElement("div");
      o.className = "orbit";
      o.style.cssText = `left:${(rnd() * 100).toFixed(2)}%;top:${(rnd() * 100).toFixed(2)}%;` +
        `animation-duration:${duration.toFixed(1)}s;animation-delay:${(-rnd() * duration).toFixed(1)}s;` +
        `animation-direction:${rnd() > .5 ? "reverse" : "normal"}`;
      const f = document.createElement("span");
      f.className = "flake";
      f.style.cssText = `width:${size.toFixed(1)}px;height:${size.toFixed(1)}px;opacity:${(.18 + rnd() * .55).toFixed(2)};` +
        `transform:translateX(${radius.toFixed(0)}px);filter:blur(${size > 4.8 ? 1.2 : 0}px)`;
      o.appendChild(f); snow.appendChild(o);
    }
    $("enter").onclick = () => { $("landing").classList.add("gone"); setTimeout(() => $("details").focus(), 700); };
  })();

  setState("idle");
</script>
</body>
</html>
"""

# Embed the logo in the page
logo_html = ""
LOGO_FILE = Path(LOGO_PATH.strip().strip('"'))
if LOGO_FILE.is_file():
    mime = mimetypes.guess_type(LOGO_FILE.name)[0] or "image/png"
    logo_html = f'<img class="logo" src="data:{mime};base64,{base64.b64encode(LOGO_FILE.read_bytes()).decode()}" alt="الشعار">'
HTML = HTML.replace("__LOGO__", logo_html)


# ---------- Start the application ----------
WINDOW = webview.create_window("الخدمات الذاتية - الاستفسارات", html=HTML, js_api=Api(),
                      width=1200, height=900, min_size=(900, 600))
webview.start()
